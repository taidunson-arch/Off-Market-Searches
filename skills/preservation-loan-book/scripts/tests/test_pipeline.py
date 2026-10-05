"""End-to-end: run_agency_pipeline.py on the fixture set produces every output; the brief quotes the calendar header verbatim;
Board Totals in the brief == Board_Totals tab == units_at_risk.json; no A/B/C, no PII; a prior run yields status_flips.csv;
the pha_am profile treats its own assets as self_owned."""
from __future__ import annotations

import json
import os
import re
import subprocess
import sys
import tempfile

import openpyxl
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from _fixture_runs import FX, PACK, SCRIPTS, by_name, hfa_fixture_run  # noqa: E402
from plb.pii import assert_no_pii  # noqa: E402

_RUN = {}


def _pipeline():
    if "r" in _RUN:
        return _RUN["r"]
    out = tempfile.mkdtemp(prefix="plb_pipe_")
    prior = os.path.join(out, "prior")
    os.makedirs(prior)
    pd.read_csv(os.path.join(FX, "prior_leads_scored_sample.csv"), dtype=str, keep_default_na=False).to_csv(os.path.join(prior, "leads_scored.csv"), index=False)
    cfg = {"market_id": "fixture-metro", "pack": PACK, "agency_profile": "hfa", "universe": "all", "geography_mode": "metro_core", "counties": ["41051", "41067", "41005"],
           "horizon_years": 10, "servicing_extract": os.path.join(FX, "agency_servicing_sample.csv"), "book_coverage": "full", "book_crosswalk": os.path.join(FX, "book_crosswalk_sample.csv"),
           "local_datasets": [os.path.join(FX, "ohcs_sample.csv"), os.path.join(FX, "reac_scores_sample.csv"), os.path.join(FX, "ohcs_forecast_sample.csv")],
           "prior_run": prior, "pii_scope": "organization", "board_packet": True, "workbook_name": "Fixture_Preservation_LoanBook_10yr.xlsx", "out_dir": out,
           "as_of_date": "2026-10-04", "pipeline_mode": True, "pipeline_root": out, "operations_db": os.path.join(out, "agency.sqlite")}
    cp = os.path.join(out, "run_config.json")
    json.dump(cfg, open(cp, "w"))
    p = subprocess.run([sys.executable, os.path.join(SCRIPTS, "run_agency_pipeline.py"), "--config", cp], capture_output=True, text=True, cwd=SCRIPTS)
    assert p.returncode == 0, p.stderr[-4000:]
    _RUN["r"] = {"root": out, "dir": json.load(open(os.path.join(out, "2026-10-04", "latest.json")))["run_dir"], "stdout": p.stdout}
    return _RUN["r"]


def test_every_output_written():
    r = _pipeline()
    d = r["dir"]
    for f in ("instruments.csv", "covenants.csv", "risk_dimensions.csv", "financial_history.json", "leads.csv", "events.csv", "rejects.csv", "merge_log.csv", "servicing_unmatched.csv", "book_join_gaps.csv", "calendar_summary.json", "leads_scored.csv",
              "units_at_risk.json", "notice_compliance_queue.csv", "nofa_targets.csv", "agency_calendar.csv", "sponsor_exposure.csv", "optout_qc_responses.csv",
              "Fixture_Preservation_LoanBook_10yr.xlsx", "Fixture_Board_Packet.xlsx", "board_packet.md", "brief.md", "status_flips.csv", "kpi_summary.json", "result.json",
              "manifest.json", "sources_used.json", "handoff/critical-dates-tracker.json", "handoff/documents_to_request.csv", "handoff/sos_worklist.csv", "handoff/board_packet.csv"):
        assert os.path.exists(os.path.join(d, f)), f
    m = json.load(open(os.path.join(d, "manifest.json")))
    assert m["flags"]["agency_profile"] == "hfa" and m["flags"]["pii_scope"] == "organization" and m["flags"]["book_coverage"] == "full" and m["plb_version"]
    assert [i["schema"] for i in m["inputs"]][0] == "agency_servicing_extract" and m["stages"]["calendar"]["header_matches_query"] is True
    assert "records_classification" in m["flags"]
    cp = json.load(open(os.path.join(r["root"], "data", "status", "fixture-metro", "asset_management", "preservation-loan-book.json")))
    assert cp["agent"] == "preservation-loan-book" and cp["phase"] == "asset_management" and cp["status"] in ("COMPLETE", "PARTIAL")
    wb = openpyxl.load_workbook(os.path.join(d, "Fixture_Preservation_LoanBook_10yr.xlsx"), read_only=True)
    assert wb.sheetnames[:3] == ["Summary", "Board_Totals", "Intervention_Queue"] and {"Book_Watchlist", "Notice_Compliance", "Agency_Calendar", "Status_Flips", "Rejects_Verify"} <= set(wb.sheetnames)
    assert set(openpyxl.load_workbook(os.path.join(d, "Fixture_Board_Packet.xlsx"), read_only=True).sheetnames) >= {"Board_Totals", "Preservation_Queue", "Redaction_Log"}


def test_hfa_operations_snapshot_and_independent_reconciliation():
    import sqlite3
    from validate_book import reconcile
    r = _pipeline()
    db = sqlite3.connect(os.path.join(r["root"], "agency.sqlite"))
    try:
        assert db.execute("SELECT count(*) FROM runs").fetchone()[0] == 1
        assert db.execute("SELECT count(*) FROM instruments").fetchone()[0] >= 10
        assert db.execute("SELECT count(*) FROM cases").fetchone()[0] > 0
        assert db.execute("SELECT count(*) FROM covenants").fetchone()[0] == db.execute("SELECT count(*) FROM instruments").fetchone()[0]
    finally:
        db.close()
    # Independent source tape count, not a count copied from generated output.
    source = pd.read_csv(os.path.join(FX, "agency_servicing_sample.csv"))
    controls = {"source": "synthetic servicing fixture (not real HFA acceptance)", "reviewer": "test",
                "totals": {"instruments": len(source), "public_upb": float(source[source["book_kind"] == "loan"]["upb"].sum())}}
    report = reconcile(r["dir"], controls)
    assert report["passed"], report
    controls["totals"]["public_upb"] += 1
    assert not reconcile(r["dir"], controls)["passed"]


def test_brief_header_verbatim_and_board_totals_agree():
    r = _pipeline()
    d = r["dir"]
    brief = open(os.path.join(d, "brief.md"), encoding="utf-8").read()
    cal = json.load(open(os.path.join(d, "calendar_summary.json")))
    uar = json.load(open(os.path.join(d, "units_at_risk.json")))
    assert cal["header"] in brief and "agency act-by" in cal["header"]
    h = uar["headline"]
    assert f"{h['properties']} properties / {h['units_at_risk']} restricted units / {h['hap_units_at_risk']} HAP units" in brief
    ws = openpyxl.load_workbook(os.path.join(d, "Fixture_Preservation_LoanBook_10yr.xlsx"), read_only=True)["Board_Totals"]
    rows = list(ws.iter_rows(values_only=True))
    tab = {}
    for row in rows:
        if row[0] in (None, "") and row[1] in uar["by_owner_cliff_band"]:
            tab[row[1]] = row[2:5]
        if row[0] and row[0] not in ("ALL",) and row[1] is None and isinstance(row[0], str) and not row[0].startswith("Board"):
            break
    for band, vals in uar["by_owner_cliff_band"].items():
        if band in tab:
            assert tab[band][0] == vals["properties"] and float(tab[band][1]) == float(vals["units_at_risk"]), band
            m = re.search(rf"\| {re.escape(band)} \| (\d+) \| ([\d.]+) \|", brief)
            assert m and int(m.group(1)) == vals["properties"], band
    assert "## Run Summary" in brief and "## Board Totals" in brief and "## Intervention Queue" in brief and "## Notice Compliance" in brief
    assert "## Status Flips Since Last Run" in brief and "## Data Gaps and Verification Queue" in brief and "bands are months" in brief.lower()
    for banned in ("Tier A", "Tier B", "Tier C", "motivated seller", "approach angle", "skip trace", "IOI"):
        assert banned not in brief and banned not in open(os.path.join(d, "board_packet.md"), encoding="utf-8").read()


def test_scored_outputs_clean_and_flips_present():
    r = _pipeline()
    d = r["dir"]
    scored = pd.read_csv(os.path.join(d, "leads_scored.csv"), dtype=str, keep_default_na=False)
    assert assert_no_pii(scored) == [] and set(scored["queue_band"]) <= {"ESCALATE", "ACT", "PLAN", "WATCH", "EXCLUDED"}
    assert "tier" not in scored.columns and "motivation_score" not in scored.columns and "outreach_template" not in scored.columns
    flips = pd.read_csv(os.path.join(d, "status_flips.csv"), dtype=str, keep_default_na=False)
    assert len(flips) > 0 and {"loan_extended", "covenant_default"} <= set(flips["flip_type"])
    res = json.load(open(os.path.join(d, "result.json")))
    assert res["run"]["skill"] == "preservation-loan-book" and res["summary"]["units_at_risk"]["headline"] == json.load(open(os.path.join(d, "units_at_risk.json")))["headline"]
    assert res["run"]["header"] and res["handoffs"]
    nc = pd.read_csv(os.path.join(d, "notice_compliance_queue.csv"), dtype=str, keep_default_na=False)
    assert len(nc) and nc["statute_cite"].str.contains("verify").all() and (nc[nc["req_id"].str.startswith("PUSH")]["window_start"] != "").any()


def test_fixture_run_named_expectations():
    r = hfa_fixture_run()
    s = r["scored"]
    vg = by_name(s, "Village Garden Apartments")
    assert vg["primary_route"] == "servicing_watch" and "notice_compliance" in vg["secondary_routes"] and vg["queue_band"] == "PLAN" and 40 <= float(vg["intervention_score"]) <= 46
    assert vg["agency_action_type"] == "PUSH_WINDOW_PREP" and "push_window_open_prep" in vg["signals"] and "RECORDS_REQUEST_DUE" not in vg["events_in_horizon"]
    ho = by_name(s, "Hollow Oak Commons")
    assert ho["universe"] == "our_book" and ho["in_inventory"].lower() == "false" and ho["primary_route"] == "servicing_watch" and float(ho["intervention_score"]) == 26.0 and ho["queue_band"] == "WATCH"
    cg = by_name(s, "Cathedral Gardens")
    assert cg["primary_route"] == "qc_admin" and cg["intervention"] in ("qc_request_acknowledgement", "qc_marketing_plan")
    g42 = by_name(s, "GOING 42")
    assert g42["primary_route"] == "none" and g42["queue_band"] == "WATCH" and g42["exclusion_reason"] == ""
    sa = by_name(s, "St Anthony Village")
    assert sa["primary_route"] == "recap_committee" and sa["queue_band"] == "ACT"
    hill = by_name(s, "Fixture Hillside Senior")
    assert hill["primary_route"] == "optout_response" and hill["agency_action_type"] == "HAP_OPTOUT_PACKAGE_DUE" and hill["book_match"] == "unmatched_expected"
    park = by_name(s, "Fixture Park Terrace")
    assert park["queue_band"] == "EXCLUDED" and r["uar"]["by_owner_cliff_band"]["stale_contract_date_verify"]["properties"] >= 1 and r["uar"]["by_owner_cliff_band"]["OVERDUE"]["hap_units_at_risk"] == 0
    albina = by_name(s, "Fixture Albina HAP Only")
    assert albina["primary_route"] == "optout_response" and albina["notice_window_state"] == "" and "NOTICE_COMPLIANCE_BREACH" not in albina["events_in_horizon"]


def test_pha_profile_self_owned_assets():
    r = hfa_fixture_run(profile="pha_am", book="pha_book_sample.csv", coverage="partial")
    y = by_name(r["scored"], "Yards at Union Station A")
    assert y["book_match"] == "self_owned" and y["universe"] == "our_book" and y["owner_type"] == "self_owned"
    assert y["primary_route"] == "recap_committee" and y["intervention"] == "pha_repositioning" and "notice_compliance" not in y["secondary_routes"]
    assert y["owner_cliff_band"] == "URGENT" and y["units_at_risk"] == "158" and r["uar"]["columns"][2] == "owned_units"
    assert "optout_response" not in set(r["scored"]["primary_route"])  # pha_am does not administer HAP opt-outs unless designee / PBCA


def test_degraded_universe_only_run_without_servicing_extract():
    """R15: no servicing extract -> every row book_absent, the brief states the degraded posture, manifest / result.json still land
    (servicing_unmatched.csv and book_join_gaps.csv are header-only, never 0-byte)."""
    out = tempfile.mkdtemp(prefix="plb_pipe_nobook_")
    cfg = {"market_id": "fixture-metro", "pack": PACK, "agency_profile": "hfa", "universe": "all", "geography_mode": "metro_core", "counties": ["41051", "41067", "41005"],
           "horizon_years": 10, "servicing_extract": None, "book_coverage": "partial",
           "local_datasets": [os.path.join(FX, "ohcs_sample.csv")], "pii_scope": "organization", "board_packet": True,
           "workbook_name": "Fixture_NoBook_10yr.xlsx", "out_dir": out, "as_of_date": "2026-10-04"}
    cp = os.path.join(out, "run_config.json")
    json.dump(cfg, open(cp, "w"))
    p = subprocess.run([sys.executable, os.path.join(SCRIPTS, "run_agency_pipeline.py"), "--config", cp], capture_output=True, text=True, cwd=SCRIPTS)
    assert p.returncode == 0, p.stderr[-4000:]
    d = json.load(open(os.path.join(out, "2026-10-04", "latest.json")))["run_dir"]
    for f in ("brief.md", "manifest.json", "result.json", "leads_scored.csv", "servicing_unmatched.csv", "book_join_gaps.csv"):
        assert os.path.exists(os.path.join(d, f)), f
        assert os.path.getsize(os.path.join(d, f)) > 0, f"{f} is 0 bytes"
    brief = open(os.path.join(d, "brief.md"), encoding="utf-8").read()
    assert "book absent: capital factor 0 for all rows" in brief
    leads = pd.read_csv(os.path.join(d, "leads_scored.csv"), dtype=str, keep_default_na=False)
    assert (leads["book_match"] == "book_absent").all() and (leads["universe"] == "universe_not_held").all()
    assert not leads["verify_flags"].str.contains("Verify — Book Join").any()
    m = json.load(open(os.path.join(d, "manifest.json")))
    assert m["flags"]["book_coverage"] == "partial" and isinstance(m["flags"]["regulatory_horizon_years"], int)
