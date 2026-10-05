"""openpyxl renderer for the preservation-loan-book working file and the board packet (references/output-contract.md Section 2-3).

Working-file tabs (order fixed): Summary, Board_Totals, Intervention_Queue, Book_Watchlist, Preservation_Queue, Sponsor_Exposure,
Notice_Compliance, Agency_Calendar, Mandate_Fit, Events, Regulatory, Owners_Sponsors, Book_Join_Gaps, Status_Flips, Sources_Vintages,
Assumptions, Rejects_Verify. Board packet: Board_Totals, Preservation_Queue (organization + registered agent + program facts),
Sponsor_Exposure, Status_Flips, Redaction_Log, Assumptions.

Cells are literals (no formulas). Queue fills ESCALATE F4CCCC / ACT FFEB9C / PLAN E7E6E6; basis fills on first_*_basis,
our_maturity_basis, senior_maturity_basis, agency_action_basis. A/B/C never appear.
"""
from __future__ import annotations

import json
import math
from typing import Any, Dict, Iterable, List, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .pii import RECORDS_CLASSIFICATION, apply_pii_scope
from .schema import AM_STATUS, ROUTES
from .units import BAND_ROWS

QUEUE_FILL = {"ESCALATE": "F4CCCC", "ACT": "FFEB9C", "PLAN": "E7E6E6"}
BASIS_FILL = {"RECORDED": "C6EFCE", "REPORTED": "DDEBF7", "DERIVED": "E2EFDA", "ESTIMATED": "FFF2CC", "PROXY": "E7E6E6"}
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
BASIS_COLS = ("first_debt_basis", "first_reg_basis", "our_maturity_basis", "senior_maturity_basis", "agency_action_basis")

QUEUE_FRONT = ["queue_band", "intervention_score", "board_impact", "primary_route", "secondary_routes", "universe", "property_name", "address", "jurisdiction",
               "units", "units_at_risk", "hap_units_at_risk", "prac_units_at_risk", "psh_units_at_risk", "owner_name", "owner_type", "programs",
               "agency_action_type", "agency_action_date", "owner_cliff_type", "owner_cliff_date", "owner_cliff_band",
               "first_reg_event_type", "first_reg_event_date", "first_reg_months_out", "first_reg_basis", "first_debt_event_type", "first_debt_event_date",
               "first_debt_months_out", "first_debt_basis", "public_upb", "our_maturity", "affordability_end", "covenant_status", "am_officer",
               "intervention", "intervention_owner", "statutory_cite", "notice_status", "verify_flags"]
BOOK_COLS = ["property_id", "property_name", "jurisdiction", "primary_route", "queue_band", "intervention_score", "book_match", "book_join_grade", "book_kind",
             "agency_loan_ids", "grant_ids", "agency_programs", "public_upb", "public_upb_at_risk", "our_rate", "payment_type", "our_maturity", "our_maturity_basis",
             "affordability_end", "recapture_type", "recapture_method", "recapture_amount", "recapture_exposure", "covenant_status", "senior_lien_type",
             "senior_maturity", "senior_maturity_basis", "senior_upb", "coterminous_senior_cliff", "am_officer", "owner_cliff_type", "owner_cliff_date", "owner_cliff_band",
             "agency_action_type", "agency_action_date", "intervention", "verify_flags"]
PRES_COLS = ["property_id", "property_name", "address", "jurisdiction", "primary_route", "secondary_routes", "queue_band", "intervention_score", "board_impact",
             "units", "units_at_risk", "hap_units_at_risk", "prac_units_at_risk", "psh_units_at_risk", "owner_name", "owner_type", "registered_agent", "programs",
             "owner_cliff_type", "owner_cliff_date", "owner_cliff_band", "withdrawal_anchor_date", "push_anchor_source", "notice_status", "recap_status",
             "mandate_fit", "mandate_eligible_products", "intervention", "intervention_owner", "statutory_cite", "verify_flags"]
WIDTHS = {"property_id": 20, "property_name": 34, "address": 30, "jurisdiction": 14, "owner_name": 34, "owner_type": 24, "programs": 24, "intervention": 30,
          "statutory_cite": 44, "verify_flags": 40, "signals": 40, "next_action": 50, "verify_before_action": 50, "events_in_horizon": 60, "am_officer": 18,
          "registered_agent": 28, "notice_address": 36, "derivation": 44, "source": 30}


def _clean(v):
    if v is None:
        return None
    if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
        return None
    if isinstance(v, (list, dict)):
        return json.dumps(v, ensure_ascii=False)
    if isinstance(v, str) and len(v) > 32000:
        return v[:32000]
    return v


def _write_df(ws, df: pd.DataFrame, widths: Optional[Dict[str, int]] = None, default_width: int = 14, freeze: bool = True):
    cols = list(df.columns)
    ws.append(cols)
    for c in ws[1]:
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for row in df.itertuples(index=False):
        ws.append([_clean(v) for v in row])
    w = dict(WIDTHS)
    w.update(widths or {})
    for i, name in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w.get(name, default_width)
    if freeze:
        ws.freeze_panes = "A2"
    if len(cols):
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(1, len(df) + 1)}"


def _fill_bases(ws):
    hdr = {c.value: i + 1 for i, c in enumerate(ws[1])}
    for r in range(2, ws.max_row + 1):
        for bc in BASIS_COLS + ("basis",):
            if bc in hdr:
                cell = ws.cell(row=r, column=hdr[bc])
                if cell.value in BASIS_FILL:
                    cell.fill = PatternFill("solid", fgColor=BASIS_FILL[cell.value])
                    if cell.value == "PROXY":
                        cell.font = Font(italic=True)
        if "queue_band" in hdr:
            cell = ws.cell(row=r, column=hdr["queue_band"])
            if cell.value in QUEUE_FILL:
                cell.fill = PatternFill("solid", fgColor=QUEUE_FILL[cell.value])


def _kv_sheet(ws, rows: Iterable[Any], title: Optional[str] = None):
    if title:
        ws.append([title])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([])
    for r in rows:
        ws.append([_clean(x) for x in (r if isinstance(r, (list, tuple)) else [r])])
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 30


def _cols(df: pd.DataFrame, want: List[str]) -> pd.DataFrame:
    return df[[c for c in want if c in df.columns]].copy()


def board_totals_rows(uar: Dict[str, Any]) -> List[List[Any]]:
    cols = uar.get("columns") or []
    rows: List[List[Any]] = [["Board totals by owner_cliff_band (months, not days) x jurisdiction", "", *cols]]
    labels = {"stale_contract_date_verify": "stale contract date — verify", "no_dated_cliff": "no dated cliff", "BEYOND": "beyond horizon"}
    rows.append(["ALL", "", *[""] * len(cols)])
    tot = {c: 0 for c in cols}
    for b in BAND_ROWS:
        v = (uar.get("by_owner_cliff_band") or {}).get(b, {})
        rows.append(["", labels.get(b, b), *[v.get(c, 0) for c in cols]])
        for c in cols:
            tot[c] = round(tot[c] + (v.get(c, 0) or 0), 1)
    rows.append(["", "TOTAL", *[tot[c] for c in cols]])
    for jur, bands in sorted((uar.get("by_jurisdiction") or {}).items()):
        rows.append([jur, "", *[""] * len(cols)])
        for b in BAND_ROWS:
            v = bands.get(b, {})
            if not v.get("properties"):
                continue
            rows.append(["", labels.get(b, b), *[v.get(c, 0) for c in cols]])
    rows.append([])
    rows.append(["Units_by_Year (owner cliff year)", "", *cols])
    for y, v in (uar.get("by_year") or {}).items():
        rows.append(["", y, *[v.get(c, 0) for c in cols]])
    rows.append([])
    h = uar.get("headline") or {}
    rows.append(["Headline (owner cliff inside 36 months, or OVERDUE without termination evidence — verify)", f"{h.get('properties', 0)} properties / {h.get('units_at_risk', 0)} restricted units / {h.get('hap_units_at_risk', 0)} HAP units / "
                                                            f"{h.get('prac_units_at_risk', 0)} PRAC units / {h.get('psh_units_at_risk', 0)} PSH units / ${h.get('public_upb_at_risk', 0):,.0f} public UPB"])
    rows.append(["public_upb_total (matched / self_owned book rows)", uar.get("public_upb_total", 0)])
    rows.append(["public_upb_on_watch", uar.get("public_upb_on_watch", 0)])
    rows.append(["book verdict", uar.get("book_verdict", ""), uar.get("book_verdict_rule", "")])
    k = uar.get("kpi") or {}
    rows.append(["KPI: units preserved since last run", k.get("units_preserved_since_prior", 0)])
    rows.append(["KPI: units lost since last run (evidence-based)", k.get("units_lost_since_prior", 0)])
    rows.append(["KPI: status flips", json.dumps(k.get("flips") or {})])
    return rows


def _board_sheet(ws, uar: Dict[str, Any]):
    rows = board_totals_rows(uar)
    for r in rows:
        ws.append([_clean(x) for x in r])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in range(2, ws.max_row + 1):
        lab = ws.cell(row=r, column=2).value
        if lab == "TOTAL":
            for c in ws[r]:
                c.font = Font(bold=True)
        if lab == "stale contract date — verify":
            for c in ws[r]:
                c.font = Font(italic=True, color="808080")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 30
    for i in range(3, 12):
        ws.column_dimensions[get_column_letter(i)].width = 16


def sponsor_exposure(leads: pd.DataFrame) -> pd.DataFrame:
    from .entities import normalize_org
    rows = []
    if not len(leads):
        return pd.DataFrame(columns=["sponsor_org", "owner_name_example", "properties", "units", "restricted_units", "hap_units", "our_upb", "upb_share", "cliffs_le_36mo", "covenant_defaults", "open_findings", "sponsor_concentration"])
    df = leads.copy()
    df["_org"] = df["owner_name"].map(normalize_org)
    df = df[df["_org"] != ""]
    total_upb = pd.to_numeric(df["public_upb"], errors="coerce").fillna(0).sum() if "public_upb" in df else 0
    for org, g in df.groupby("_org"):
        upb = pd.to_numeric(g["public_upb"], errors="coerce").fillna(0).sum() if "public_upb" in g else 0
        cliffs = int(g["owner_cliff_band"].isin(["OVERDUE", "CRITICAL", "URGENT", "APPROACHING"]).sum()) if "owner_cliff_band" in g else 0
        defaults = int((g.get("covenant_status", pd.Series(dtype=str)) == "default").sum())
        findings = int(g["signals"].astype(str).str.contains("noncompliance_finding", na=False).sum()) if "signals" in g else 0
        share = round(upb / total_upb, 3) if total_upb else 0.0
        rows.append({"sponsor_org": org, "owner_name_example": g["owner_name"].iloc[0], "properties": int(len(g)),
                     "units": int(pd.to_numeric(g["units"], errors="coerce").fillna(0).sum()),
                     "restricted_units": int(pd.to_numeric(g["restricted_units"], errors="coerce").fillna(0).sum()) if "restricted_units" in g else 0,
                     "hap_units": int(pd.to_numeric(g["hap_units_at_risk"], errors="coerce").fillna(0).sum()) if "hap_units_at_risk" in g else 0,
                     "our_upb": round(upb, 0), "upb_share": share, "cliffs_le_36mo": cliffs, "covenant_defaults": defaults, "open_findings": findings,
                     "sponsor_concentration": bool(share >= 0.10 or cliffs >= 3)})
    out = pd.DataFrame(rows)
    return out.sort_values(["cliffs_le_36mo", "our_upb"], ascending=[False, False]).reset_index(drop=True) if len(out) else out


def _summary_rows(leads: pd.DataFrame, calendar: Dict[str, Any], run_meta: Dict[str, Any], sources, uar: Dict[str, Any]) -> List[Any]:
    rows: List[Any] = [
        ["agency_profile", run_meta.get("agency_profile", "")], ["agency_name", run_meta.get("agency_name", "")],
        ["as_of_date", run_meta.get("as_of_date", "")], ["geography_mode", run_meta.get("geography_mode", "")],
        ["counties (FIPS)", ", ".join(run_meta.get("counties", []) or [])], ["horizon_years", run_meta.get("horizon_years", "")],
        ["universe filter", run_meta.get("universe", "all")], ["book_coverage", run_meta.get("book_coverage", "")], ["pii_scope", run_meta.get("pii_scope", "organization")],
        ["records_classification", RECORDS_CLASSIFICATION],
        ["rows in geography (pre status filter)", (calendar or {}).get("rows_in_geography", "")], ["rows excluded In Development", (calendar or {}).get("rows_excluded_in_development", "")],
        ["leads", len(leads)], [],
        ["Basis breakdown header (quote this before any count; bands are months, not days)", (calendar or {}).get("header", "")],
    ]
    for b in ("RECORDED", "REPORTED", "DERIVED", "ESTIMATED", "PROXY"):
        rows.append([b, (calendar or {}).get("by_basis", {}).get(b, 0)])
    rows += [["Suppressed", (calendar or {}).get("suppressed", 0)], ["Rejected", (calendar or {}).get("rejected", 0)],
             ["helper deadlines", (calendar or {}).get("helper_events", 0)], ["agency act-by", (calendar or {}).get("agency_act_by", 0)],
             ["agency act-by next 90 days", (calendar or {}).get("agency_act_by_next_90_days", 0)], ["stale contract dates", (calendar or {}).get("stale_contract_dates", 0)], []]
    if "queue_band" in leads:
        rows.append(["Queue counts"])
        for t, n in leads["queue_band"].value_counts().items():
            rows.append([t, int(n)])
        rows.append([])
    if "primary_route" in leads:
        rows.append(["Primary route counts"])
        for t, n in leads["primary_route"].value_counts().items():
            rows.append([t, int(n)])
        rows.append([])
    if "book_match" in leads:
        rows.append(["book_match counts", json.dumps(leads["book_match"].value_counts().to_dict())])
        rows.append(["book_join_grade counts", json.dumps(leads[leads["book_join_grade"].astype(str) != ""]["book_join_grade"].value_counts().to_dict()) if "book_join_grade" in leads else ""])
    rows.append(["book verdict", uar.get("book_verdict", "")])
    unverified = [s.get("source_id") for s in (sources or []) if not str(s.get("url_verified_live", s.get("verified_live", "false"))).lower().startswith("true")]
    rows.append(["Sources still verified_live=false (URL not fetched in this build)", "; ".join(str(x) for x in unverified) if unverified else "none listed"])
    rows.append(["Degraded-run note", run_meta.get("degraded_note", "")])
    return rows


def build_workbook(out_path: str, leads: pd.DataFrame, events: pd.DataFrame, calendar: Dict[str, Any], rejects: Optional[pd.DataFrame] = None,
                   assumptions: Optional[Dict[str, Any]] = None, sources: Optional[List[Dict[str, Any]]] = None, status_flips: Optional[pd.DataFrame] = None,
                   run_meta: Optional[Dict[str, Any]] = None, units_at_risk: Optional[Dict[str, Any]] = None, agency_calendar: Optional[pd.DataFrame] = None,
                   notice_compliance: Optional[pd.DataFrame] = None, mandate_fit: Optional[pd.DataFrame] = None, book_join_gaps: Optional[pd.DataFrame] = None,
                   pii_scope: str = "organization") -> str:
    wb = Workbook()
    run_meta = run_meta or {}
    uar = units_at_risk or {}
    leads, _ = apply_pii_scope(leads.copy(), pii_scope)
    events = events.copy() if events is not None else pd.DataFrame()

    ws = wb.active
    ws.title = "Summary"
    _kv_sheet(ws, _summary_rows(leads, calendar, run_meta, sources, uar), title="preservation-loan-book working file")

    ws = wb.create_sheet("Board_Totals")
    _board_sheet(ws, uar)

    ws = wb.create_sheet("Intervention_Queue")
    q = leads[leads["primary_route"].isin(ROUTES)] if "primary_route" in leads else leads
    cols = [c for c in QUEUE_FRONT if c in q.columns] + [c for c in q.columns if c not in QUEUE_FRONT]
    t = q[cols].copy()
    t.insert(0, "AM status", "New")
    _write_df(ws, t)
    _fill_bases(ws)
    if ws.max_row >= 2:
        dv = DataValidation(type="list", formula1='"' + ",".join(AM_STATUS) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"A2:A{ws.max_row}")

    ws = wb.create_sheet("Book_Watchlist")
    bw = leads[leads["universe"] == "our_book"] if "universe" in leads else leads.iloc[0:0]
    t = _cols(bw, BOOK_COLS)
    t.insert(0, "AM status", ["Monitoring" if r == "none" else "New" for r in bw["primary_route"]] if "primary_route" in bw else "Monitoring")
    _write_df(ws, t)
    _fill_bases(ws)
    if ws.max_row >= 2:
        dv = DataValidation(type="list", formula1='"' + ",".join(AM_STATUS + ["Monitoring"]) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"A2:A{ws.max_row}")

    ws = wb.create_sheet("Preservation_Queue")
    pq = leads[leads["universe"] == "universe_not_held"] if "universe" in leads else leads.iloc[0:0]
    _write_df(ws, _cols(pq, PRES_COLS))
    _fill_bases(ws)

    ws = wb.create_sheet("Sponsor_Exposure")
    _write_df(ws, sponsor_exposure(leads), {"sponsor_org": 40, "owner_name_example": 36})

    ws = wb.create_sheet("Notice_Compliance")
    nc = notice_compliance if notice_compliance is not None else pd.DataFrame(columns=["req_id", "property_id", "property_name", "statute_cite", "requirement", "window_start", "window_end",
                                                                                      "due_date", "status", "evidence", "finding", "remediation_ask", "owner_response", "date_resolved"])
    _write_df(ws, nc, {"requirement": 40, "statute_cite": 36, "finding": 40, "remediation_ask": 40, "evidence": 36})

    ws = wb.create_sheet("Agency_Calendar")
    ac = agency_calendar if agency_calendar is not None else pd.DataFrame(columns=["property_id", "property_name", "agency_action_type", "due_date", "months_out", "urgency", "basis",
                                                                                 "derived_from", "agency_owner", "statutory_cite", "status"])
    _write_df(ws, ac, {"agency_action_type": 30, "derived_from": 40, "statutory_cite": 40})
    _fill_bases(ws)

    ws = wb.create_sheet("Mandate_Fit")
    mf = mandate_fit if mandate_fit is not None else _cols(leads, ["property_id", "property_name", "programs", "owner_type", "units", "mandate_fit", "mandate_eligible_products", "mandate_ineligible_reason", "primary_route"])
    _write_df(ws, mf, {"mandate_eligible_products": 36, "mandate_ineligible_reason": 60})

    ws = wb.create_sheet("Events")
    if len(events):
        _write_df(ws, events, {"event_type": 28, "verify_flag": 26, "alt_dates": 24, "detail": 30})
        _fill_bases(ws)
    else:
        ws.append(["no events"])

    ws = wb.create_sheet("Regulatory")
    reg_cols = ["property_id", "property_name", "programs", "hud_contract", "ami_30_60_units", "ami_80_units", "market_rate_units", "rental_assistance_units",
                "first_reg_event_type", "first_reg_event_date", "first_reg_months_out", "first_reg_basis", "first_reg_source", "withdrawal_anchor_date", "push_anchor_source",
                "notice_window_state", "notice_status", "tenant_notice_status", "hap_renewal_request_status", "next_expected_expiration", "qc_status", "qc_waived", "rofr_recorded", "verify_flags"]
    _write_df(ws, _cols(leads, reg_cols))
    _fill_bases(ws)

    ws = wb.create_sheet("Owners_Sponsors")
    o_cols = ["property_id", "property_name", "owner_name", "owner_type", "owner_archetype", "developer_name", "manager_name", "registered_agent", "registered_agent_address",
              "sos_status", "notice_address", "notice_address_source", "sponsor_contact_role", "org_resolution_grade", "sponsor_cliff_count", "primary_route", "intervention", "intervention_owner"]
    _write_df(ws, _cols(leads, o_cols))

    ws = wb.create_sheet("Book_Join_Gaps")
    bj = book_join_gaps if book_join_gaps is not None else pd.DataFrame(columns=["property_id", "property_name", "book_match", "expected_book_flag", "note"])
    _write_df(ws, bj, {"note": 50})

    ws = wb.create_sheet("Status_Flips")
    if status_flips is not None and len(status_flips):
        _write_df(ws, status_flips, {"evidence": 50})
    else:
        ws.append(["no prior run supplied"])

    ws = wb.create_sheet("Sources_Vintages")
    src_df = pd.DataFrame(sources or [])
    want = ["source_id", "file", "sha256", "vintage", "vintage_source", "file_verified", "url_verified_live", "verified_live", "access", "url", "note"]
    if len(src_df) == 0:
        src_df = pd.DataFrame(columns=want)
    else:
        for c in want:
            if c not in src_df.columns:
                src_df[c] = ""
        src_df = src_df[want + [c for c in src_df.columns if c not in want]]
    _write_df(ws, src_df, {"source_id": 30, "file": 40, "sha256": 24, "url": 50, "file_verified": 26, "url_verified_live": 26, "note": 50})
    hdr = {c.value: i + 1 for i, c in enumerate(ws[1])}
    for r in range(2, ws.max_row + 1):
        for col in ("url_verified_live", "verified_live"):
            if col in hdr and not str(ws.cell(row=r, column=hdr[col]).value or "").lower().startswith("true"):
                ws.cell(row=r, column=hdr[col]).fill = PatternFill("solid", fgColor="FFEB9C")

    ws = wb.create_sheet("Assumptions")
    a_rows = [["parameter", "value", "as_of", "source"]]
    for k, v in (assumptions or {}).items():
        if k.startswith("_"):
            continue
        if isinstance(v, dict) and "value" in v:
            a_rows.append([k, json.dumps(v["value"]) if isinstance(v["value"], (dict, list)) else v["value"], v.get("as_of", ""), str(v.get("source", "")) + (" [verify]" if v.get("verify") else "")])
        else:
            a_rows.append([k, json.dumps(v) if isinstance(v, (dict, list)) else v, "", ""])
    a_rows += [[], ["Scoring tables", run_meta.get("scoring_source", "references/scoring/affordable_public_am.json + shared_adjustments.json")],
               ["Basis multipliers", "RECORDED 1.00 / REPORTED 0.90 / DERIVED 0.80 / ESTIMATED 0.50 / PROXY 0.30; Verify — Ambiguous 0.40; HAP confidence 0.70 / 0.90 on that band only"],
               ["UPB-at-risk rule", "upb where any owner-cliff PRESSURE event <= 36 mo OR covenant_status in {watch, default} OR coterminous_senior_cliff"],
               ["Band rules", "Board_Totals and Units_by_Year band on owner_cliff_band; the queue sorts on action_band = min(agency act-by, owner cliff); bands are months, not days"],
               ["Equity overlay", run_meta.get("equity_overlay", "absent; stacking_geography_equity modifiers scored 0")],
               ["Statutory text status", "every cite verified_live false in this build; confirm current text before sending"],
               ["records_classification", RECORDS_CLASSIFICATION],
               ["Measure 50", "Assessed Value is never used as market value in Oregon"]]
    _kv_sheet(ws, a_rows)
    for c in ws[1]:
        c.font = Font(bold=True)

    ws = wb.create_sheet("Rejects_Verify")
    rj = rejects if rejects is not None else pd.DataFrame(columns=["property_key", "column", "raw_value", "reason", "flag", "alt_dates"])
    _write_df(ws, rj, {"property_key": 40, "column": 26, "raw_value": 24, "reason": 50, "flag": 28, "alt_dates": 24})

    wb.save(out_path)
    return out_path


def build_board_packet(out_path: str, leads: pd.DataFrame, units_at_risk: Dict[str, Any], status_flips: Optional[pd.DataFrame], assumptions: Optional[Dict[str, Any]],
                       run_meta: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    """Second workbook in pii_scope public_packet. Returns {path, redaction_log}."""
    wb = Workbook()
    packet, log = apply_pii_scope(leads.copy(), "public_packet")
    ws = wb.active
    ws.title = "Board_Totals"
    _board_sheet(ws, units_at_risk or {})
    ws = wb.create_sheet("Preservation_Queue")
    pq = packet[packet["universe"] == "universe_not_held"] if "universe" in packet else packet
    _write_df(ws, _cols(pq, [c for c in PRES_COLS if c in packet.columns]))
    _fill_bases(ws)
    ws = wb.create_sheet("Sponsor_Exposure")
    _write_df(ws, sponsor_exposure(packet), {"sponsor_org": 40})
    ws = wb.create_sheet("Status_Flips")
    if status_flips is not None and len(status_flips):
        sf = status_flips[[c for c in status_flips.columns if c in ("property_id", "property_name", "flip_type", "prior_value", "new_value", "basis", "change")]]
        _write_df(ws, sf)
    else:
        ws.append(["no prior run supplied"])
    ws = wb.create_sheet("Redaction_Log")
    _write_df(ws, pd.DataFrame(log, columns=["field", "rows_redacted", "action", "cite"]), {"field": 60, "action": 44, "cite": 60})
    ws = wb.create_sheet("Assumptions")
    rows = [["parameter", "value"], ["pii_scope", "public_packet"], ["records_classification", RECORDS_CLASSIFICATION],
            ["statutory text", "every cite verified_live false in this build; confirm current text before relying on it"],
            ["band rules", "owner_cliff_band drives Board_Totals; months, not days"], ["agency_profile", (run_meta or {}).get("agency_profile", "")]]
    for k, v in (assumptions or {}).items():
        if k.startswith("_"):
            continue
        rows.append([k, json.dumps(v["value"]) if isinstance(v, dict) and "value" in v and isinstance(v["value"], (dict, list)) else (v.get("value") if isinstance(v, dict) and "value" in v else (json.dumps(v) if isinstance(v, (dict, list)) else v))])
    _kv_sheet(ws, rows)
    wb.save(out_path)
    return {"path": out_path, "redaction_log": log}
