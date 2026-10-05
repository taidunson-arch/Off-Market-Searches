"""openpyxl renderer for the off-market-deal-finder v2 target-list workbook.

Tabs (references/output-contract.md): Summary, Targets, Events, Capital_Stack, Regulatory,
Owners_DecisionMakers, Signals, Sources_Vintages, Assumptions, Rejects_Verify, Diff_vs_Prior.

Targets: freeze A2, autofilter, tier fills (A green C6EFCE, B amber FFEB9C, C grey E7E6E6),
basis fills on the first_*_basis columns (RECORDED C6EFCE, REPORTED DDEBF7, DERIVED E2EFDA,
ESTIMATED FFF2CC, PROXY E7E6E6 italic), `Analyst status` dropdown {New, Contacted, Verifying,
Nurture, Dead}. Values are written as literals (no formulas) so recalc is optional; callers may
still run /mnt/skills/public/xlsx/scripts/recalc.py when available.
"""
from __future__ import annotations

import json
import math
from typing import Any, Dict, Iterable, List, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

TIER_FILL = {"A": "C6EFCE", "B": "FFEB9C", "C": "E7E6E6"}
BASIS_FILL = {"RECORDED": "C6EFCE", "REPORTED": "DDEBF7", "DERIVED": "E2EFDA", "ESTIMATED": "FFF2CC", "PROXY": "E7E6E6"}
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
ANALYST_STATUS = ["New", "Contacted", "Verifying", "Nurture", "Dead"]

TARGET_WIDTHS = {"property_id": 20, "property_name": 34, "address": 30, "city": 14, "zip": 8, "county_name": 12,
                 "owner_name": 34, "owner_type": 24, "programs": 26, "first_debt_event_type": 22, "first_debt_event_date": 13,
                 "first_reg_event_type": 24, "first_reg_event_date": 13, "signals": 40, "verify_flags": 28,
                 "outreach_angle": 40, "verify_before_outreach": 50, "next_action": 50, "events_in_horizon": 60,
                 "decision_maker_name": 24, "decision_maker_role": 24}

TARGET_FRONT = ["tier", "motivation_score", "route", "property_name", "address", "city", "county_name", "units", "year_built",
                "rehab_year", "owner_name", "owner_type", "programs", "first_debt_event_type", "first_debt_event_date",
                "first_debt_months_out", "first_debt_basis", "first_reg_event_type", "first_reg_event_date",
                "first_reg_months_out", "first_reg_basis", "decision_maker_name", "decision_maker_role", "contact_grade",
                "outreach_template", "compliance_gates", "verify_before_outreach", "next_action"]


def _clean(v):
    if v is None:
        return None
    if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
        return None
    if isinstance(v, (list, dict)):
        return json.dumps(v, ensure_ascii=False)
    if isinstance(v, str) and len(v) > 32000:
        return v[:32000]
    return v


def _write_df(ws, df: pd.DataFrame, widths: Optional[Dict[str, int]] = None, default_width: int = 14, freeze: bool = True):
    cols = list(df.columns)
    ws.append(cols)
    for c in ws[1]:
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for row in df.itertuples(index=False):
        ws.append([_clean(v) for v in row])
    for i, name in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(name, default_width)
    if freeze:
        ws.freeze_panes = "A2"
    if len(cols):
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(1, len(df) + 1)}"


def _kv_sheet(ws, rows: Iterable[Any], title: Optional[str] = None):
    if title:
        ws.append([title])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([])
    for r in rows:
        ws.append([_clean(x) for x in (r if isinstance(r, (list, tuple)) else [r])])
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 30


def build_workbook(out_path: str, leads: pd.DataFrame, events: pd.DataFrame, calendar: Dict[str, Any],
                   rejects: Optional[pd.DataFrame] = None, assumptions: Optional[Dict[str, Any]] = None,
                   sources: Optional[List[Dict[str, Any]]] = None, diff: Optional[pd.DataFrame] = None,
                   run_meta: Optional[Dict[str, Any]] = None) -> str:
    wb = Workbook()
    leads = leads.copy()
    events = events.copy() if events is not None else pd.DataFrame()
    run_meta = run_meta or {}

    # ------------------------------------------------------------- Summary
    ws = wb.active
    ws.title = "Summary"
    by_basis = (calendar or {}).get("by_basis", {})
    rows: List[Any] = [
        ["as_of_date", run_meta.get("as_of_date", leads["as_of_date"].iloc[0] if "as_of_date" in leads and len(leads) else "")],
        ["geography_mode", run_meta.get("geography_mode", "")],
        ["counties (FIPS)", ", ".join(run_meta.get("counties", []) or [])],
        ["asset_class", run_meta.get("asset_class", "")],
        ["horizon_years (debt)", run_meta.get("horizon_years", "")],
        ["regulatory_horizon_years", run_meta.get("regulatory_horizon_years", "")],
        ["rows in geography (pre status filter)", (calendar or {}).get("rows_in_geography", "")],
        ["rows excluded In Development", (calendar or {}).get("rows_excluded_in_development", "")],
        ["leads in list", len(leads)],
        [],
        ["Basis breakdown of dated events (quote this before the count)"],
    ]
    for b in ("RECORDED", "REPORTED", "DERIVED", "ESTIMATED", "PROXY"):
        rows.append([b, by_basis.get(b, 0)])
    rows.append(["Suppressed", (calendar or {}).get("suppressed", 0)])
    rows.append(["Rejected", (calendar or {}).get("rejected", len(rejects) if rejects is not None else 0)])
    rows.append([])
    if "tier" in leads:
        rows.append(["Tier counts"])
        for t, n in leads["tier"].value_counts().items():
            rows.append([t, int(n)])
        rows.append([])
    if "route" in leads:
        rows.append(["Route counts"])
        for t, n in leads["route"].value_counts().items():
            rows.append([t, int(n)])
        rows.append([])
    # `verified_live` is about the URL / portal. A file opened on disk is `file_verified`; it does not make the source live-verified.
    unverified = [s.get("source_id") for s in (sources or [])
                  if not str(s.get("url_verified_live", s.get("verified_live", "false"))).lower().startswith("true")]
    rows.append(["Sources still verified_live=false (URL not fetched in this build)", "; ".join(str(x) for x in unverified) if unverified else "none listed"])
    rows.append(["query_type / buyer_profile", f"{run_meta.get('query_type', '')} / {run_meta.get('buyer_profile', '')}"])
    if "query_hit" in leads and run_meta.get("query_type") not in (None, "", "all"):
        hits = int(leads["query_hit"].astype(str).str.lower().isin(["true", "1"]).sum())
        rows.append([f"{run_meta.get('query_type')} hits", f"{hits} leads carry a dated {run_meta.get('query_type')} event; {len(leads) - hits} shown as secondary"])
    rows.append(["Degraded-run note", run_meta.get("degraded_note", "")])
    _kv_sheet(ws, rows, title="off-market-deal-finder v2 target list")

    # ------------------------------------------------------------- Targets
    ws = wb.create_sheet("Targets")
    cols = [c for c in TARGET_FRONT if c in leads.columns] + [c for c in leads.columns if c not in TARGET_FRONT]
    t = leads[cols].copy()
    t.insert(0, "Analyst status", "New")
    _write_df(ws, t, TARGET_WIDTHS)
    hdr = {c.value: i + 1 for i, c in enumerate(ws[1])}
    for r in range(2, ws.max_row + 1):
        tier = ws.cell(row=r, column=hdr["tier"]).value if "tier" in hdr else None
        if tier in TIER_FILL:
            ws.cell(row=r, column=hdr["tier"]).fill = PatternFill("solid", fgColor=TIER_FILL[tier])
        for bc in ("first_debt_basis", "first_reg_basis", "balance_basis"):
            if bc in hdr:
                cell = ws.cell(row=r, column=hdr[bc])
                if cell.value in BASIS_FILL:
                    cell.fill = PatternFill("solid", fgColor=BASIS_FILL[cell.value])
                    if cell.value == "PROXY":
                        cell.font = Font(italic=True)
    if ws.max_row >= 2:
        dv = DataValidation(type="list", formula1='"' + ",".join(ANALYST_STATUS) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"A2:A{ws.max_row}")

    # ------------------------------------------------------------- Events
    ws = wb.create_sheet("Events")
    if len(events):
        _write_df(ws, events, {"property_id": 20, "event_type": 26, "source": 36, "derivation": 40, "verify_flag": 26, "alt_dates": 24})
        hdr = {c.value: i + 1 for i, c in enumerate(ws[1])}
        if "basis" in hdr:
            for r in range(2, ws.max_row + 1):
                cell = ws.cell(row=r, column=hdr["basis"])
                if cell.value in BASIS_FILL:
                    cell.fill = PatternFill("solid", fgColor=BASIS_FILL[cell.value])
                    if cell.value == "PROXY":
                        cell.font = Font(italic=True)
    else:
        ws.append(["no events"])

    # ------------------------------------------------------------- Capital_Stack
    ws = wb.create_sheet("Capital_Stack")
    cs_cols = [c for c in ["property_id", "property_name", "units", "first_debt_event_type", "first_debt_event_date", "first_debt_months_out",
                           "first_debt_basis", "first_debt_source", "est_value", "value_source", "est_noi", "noi_source", "est_loan_balance",
                           "balance_basis", "est_ltv", "est_dscr_refi", "refi_gap_pct", "equity_cushion_pct", "assumable_debt"] if c in leads.columns]
    _write_df(ws, leads[cs_cols], {"property_name": 34, "first_debt_event_type": 22, "first_debt_source": 30})

    # ------------------------------------------------------------- Regulatory
    ws = wb.create_sheet("Regulatory")
    reg_cols = [c for c in ["property_id", "property_name", "programs", "hud_contract", "ami_30_60_units", "ami_80_units", "market_rate_units",
                            "rental_assistance_units", "first_reg_event_type", "first_reg_event_date", "first_reg_months_out", "first_reg_basis",
                            "first_reg_source", "verify_flags"] if c in leads.columns]
    _write_df(ws, leads[reg_cols], {"property_name": 34, "programs": 30, "first_reg_event_type": 24, "first_reg_source": 30, "verify_flags": 30})

    # ------------------------------------------------------------- Owners_DecisionMakers
    ws = wb.create_sheet("Owners_DecisionMakers")
    o_cols = [c for c in ["property_id", "property_name", "owner_name", "owner_type", "owner_archetype", "developer_name", "manager_name",
                          "decision_maker_name", "decision_maker_role", "dm_source", "contact_grade", "route", "outreach_template",
                          "outreach_angle", "compliance_gates"] if c in leads.columns]
    _write_df(ws, leads[o_cols], {"property_name": 34, "owner_name": 36, "developer_name": 30, "manager_name": 30, "outreach_angle": 44})

    # ------------------------------------------------------------- Signals
    ws = wb.create_sheet("Signals")
    sig_rows = []
    for _, r in leads.iterrows():
        for s in str(r.get("signals") or "").split(";"):
            if s.strip():
                sig_rows.append({"property_id": r.get("property_id"), "property_name": r.get("property_name"), "signal": s.strip(), "tier": r.get("tier")})
    _write_df(ws, pd.DataFrame(sig_rows, columns=["property_id", "property_name", "signal", "tier"]), {"property_name": 34, "signal": 36})

    # ------------------------------------------------------------- Sources_Vintages
    ws = wb.create_sheet("Sources_Vintages")
    src_df = pd.DataFrame(sources or [])
    want = ["source_id", "file", "sha256", "vintage", "vintage_source", "file_verified", "url_verified_live", "verified_live", "access", "url", "note"]
    if len(src_df) == 0:
        src_df = pd.DataFrame(columns=want)
    else:
        for c in want:
            if c not in src_df.columns:
                src_df[c] = ""
        src_df = src_df[want + [c for c in src_df.columns if c not in want]]
    _write_df(ws, src_df, {"source_id": 30, "file": 40, "sha256": 24, "url": 50, "file_verified": 26, "url_verified_live": 26, "note": 50})
    hdr = {c.value: i + 1 for i, c in enumerate(ws[1])}
    for r in range(2, ws.max_row + 1):
        for col in ("url_verified_live", "verified_live"):
            if col in hdr and not str(ws.cell(row=r, column=hdr[col]).value or "").lower().startswith("true"):
                ws.cell(row=r, column=hdr[col]).fill = PatternFill("solid", fgColor="FFEB9C")

    # ------------------------------------------------------------- Assumptions
    ws = wb.create_sheet("Assumptions")
    a_rows = [["parameter", "value", "as_of", "source"]]
    for k, v in (assumptions or {}).items():
        if k.startswith("_"):
            continue
        if isinstance(v, dict) and "value" in v:
            a_rows.append([k, json.dumps(v["value"]) if isinstance(v["value"], (dict, list)) else v["value"], v.get("as_of", ""), v.get("source", "") + (" [verify]" if v.get("verify") else "")])
        else:
            a_rows.append([k, json.dumps(v) if isinstance(v, (dict, list)) else v, "", ""])
    a_rows.append([])
    a_rows.append(["Scoring tables", run_meta.get("scoring_source", "references/scoring/*.json")])
    a_rows.append(["Basis multipliers", "RECORDED 1.00 / REPORTED 0.90 / DERIVED 0.80 / ESTIMATED 0.50 / PROXY 0.30; Verify — Ambiguous 0.40"])
    a_rows.append(["Measure 50", "Assessed Value is never used as market value in Oregon; RMV or income proxy only"])
    _kv_sheet(ws, a_rows)
    for c in ws[1]:
        c.font = Font(bold=True)

    # ------------------------------------------------------------- Rejects_Verify
    ws = wb.create_sheet("Rejects_Verify")
    rj = rejects if rejects is not None else pd.DataFrame(columns=["property_key", "column", "raw_value", "reason", "flag", "alt_dates"])
    _write_df(ws, rj, {"property_key": 40, "column": 26, "raw_value": 24, "reason": 50, "flag": 28, "alt_dates": 24})

    # ------------------------------------------------------------- Diff_vs_Prior
    ws = wb.create_sheet("Diff_vs_Prior")
    if diff is not None and len(diff):
        _write_df(ws, diff, {"property_name": 34, "change": 14, "detail": 50})
    else:
        ws.append(["no prior run supplied"])

    wb.save(out_path)
    return out_path
